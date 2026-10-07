"""Reconcile official registry rows and select complete retained IGN municipality features."""

import csv
import io
from collections.abc import Iterable

from openpyxl import load_workbook
from pydantic import TypeAdapter

from spain_housing.models.dataset import (
    BoundaryCollection,
    BoundaryFeature,
    BuildRecipe,
    LocationRecord,
)
from spain_housing.storage.release import strict_json, validate_geometry


def _code_column(value: object, width: int) -> str:
    if isinstance(value, bool) or not isinstance(value, (int, str)):
        raise ValueError("INE code columns must be integers or documented digit strings")
    text = str(value).strip()
    if not text.isascii() or not text.isdigit() or len(text) > width:
        raise ValueError("invalid INE code column width/value")
    return text.zfill(width)


def _registry_table(rows: Iterable[tuple[object, ...]]) -> dict[str, tuple[str, str]]:
    iterator = iter(rows)
    header = None
    for _ in range(10):
        row = next(iterator, ())
        names = [str(value).strip() for value in row]
        if {"CPRO", "CMUN", "DC", "NOMBRE"} <= set(names):
            if len(names) != len(set(names)):
                raise ValueError("duplicate registry headers")
            header = names
            break
    if header is None:
        raise ValueError("official registry requires CPRO, CMUN, DC, NOMBRE columns")
    indexes = [header.index(name) for name in ("CPRO", "CMUN", "DC", "NOMBRE")]
    result: dict[str, tuple[str, str]] = {}
    for row in iterator:
        if not row or all(value is None or value == "" for value in row):
            continue
        if len(row) <= max(indexes):
            raise ValueError("incomplete registry row")
        province, municipality, control, name = (row[index] for index in indexes)
        code = _code_column(province, 2) + _code_column(municipality, 3)
        if code in result or not isinstance(name, str) or not name.strip():
            raise ValueError("duplicate code or invalid official registry name")
        result[code] = (name.strip(), _code_column(control, 1))
    return result


def verify_registry(data: bytes, format: str, locations: tuple[LocationRecord, ...]) -> None:
    if format == "xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            registry = (
                _registry_table(workbook.active.iter_rows(values_only=True))
                if workbook.active is not None
                else {}
            )
        finally:
            workbook.close()
    elif format == "csv":
        registry = _registry_table(
            tuple(
                tuple(row) for row in csv.reader(io.StringIO(data.decode("utf-8-sig")), strict=True)
            )
        )
    else:
        raise ValueError("registry capture must be csv or xlsx")
    for location in locations:
        if registry.get(location.location.code) != (location.official_name, location.control_digit):
            raise ValueError(
                f"official INE identity/control digit mismatch: {location.location.code}"
            )


def select_boundaries(recipe: BuildRecipe, captured_bytes: dict[str, bytes]) -> BoundaryCollection:
    selected = {loc.geometry_feature_id: loc for loc in recipe.locations}
    features: list[BoundaryFeature] = []
    for capture_id in recipe.geography.boundary_capture_ids:
        collection = TypeAdapter(dict[str, object]).validate_python(
            strict_json(captured_bytes[capture_id])
        )
        if collection.get("type") != "FeatureCollection" or collection.get("crs") not in (
            None,
            "OGC:CRS84",
        ):
            raise ValueError("source must be an OGC:CRS84 GeoJSON FeatureCollection")
        raw_features = TypeAdapter(list[dict[str, object]]).validate_python(
            collection.get("features")
        )
        for name in ("numberMatched", "numberReturned"):
            if name in collection and collection[name] != len(raw_features):
                raise ValueError("incomplete/paginated geographic response")
        links = TypeAdapter(list[dict[str, object]]).validate_python(collection.get("links", []))
        if any(link.get("rel") == "next" for link in links):
            raise ValueError("geographic capture has unretained pagination")
        for raw in raw_features:
            feature_id = str(raw.get("id", ""))
            location = selected.get(feature_id)
            if location is None:
                continue
            properties = TypeAdapter(dict[str, object]).validate_python(raw.get("properties"))
            if raw.get("type") != "Feature" or (
                properties.get("nationallevelname"),
                properties.get("nationalcode"),
                properties.get("nameunit"),
            ) != ("Municipio", location.geometry_national_code, location.official_name):
                raise ValueError("IGN feature is not the verified municipality/crosswalk")
            features.append(
                BoundaryFeature.model_validate(
                    {
                        "id": feature_id,
                        "properties": {
                            "location_code": location.location.code,
                            "boundary_version": location.boundary_version,
                        },
                        "geometry": raw.get("geometry"),
                    }
                )
            )
    boundaries = BoundaryCollection(features=tuple(features))
    validate_geometry(boundaries, recipe.locations)
    return boundaries
